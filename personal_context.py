from __future__ import annotations

import hashlib
import json
import sqlite3
import subprocess
import sys
import zipfile
import re
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path
from uuid import uuid4

import config


def read_document(filename: str) -> list[dict]:
    path = Path(filename).resolve(strict=True)
    if not path.is_file() or path.stat().st_size > 20 * 1024 * 1024:
        raise ValueError("请选择不超过 20 MB 的文件")
    suffix = path.suffix.lower()
    if suffix in (".docx", ".xlsx"):
        with zipfile.ZipFile(path) as archive:
            if len(archive.infolist()) > 5000 or sum(info.file_size for info in archive.infolist()) > 64 * 1024 * 1024:
                raise ValueError("Office 文件解压后过大，请先拆分")
    pages = []
    if suffix == ".pdf":
        import fitz
        with fitz.open(path) as document:
            if len(document) > 300:
                raise ValueError("请先拆分超过 300 页的 PDF")
            used = 0
            for index, page in enumerate(document):
                text = page.get_text()
                used += len(text)
                if used > 1000000:
                    raise ValueError("PDF 文本过长，请拆分后读取")
                pages.append((index + 1, text))
    elif suffix == ".docx":
        from docx import Document
        document = Document(path)
        text = "\n".join(p.text for p in document.paragraphs)
        text += "\n" + "\n".join(" | ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows)
        pages = [(None, text)]
    elif suffix == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in workbook:
                rows = []
                for index, row in enumerate(sheet.iter_rows(values_only=True)):
                    if index >= 20000:
                        raise ValueError("表格过大，请先导出需要的部分")
                    rows.append(" | ".join(str(value) if value is not None else "" for value in row[:100]))
                pages.append((sheet.title, "\n".join(rows)))
        finally:
            workbook.close()
    elif suffix in (".png", ".jpg", ".jpeg", ".webp", ".bmp"):
        pages = [(None, _windows_ocr(path))]
    elif suffix in {".txt", ".md", ".py", ".rs", ".js", ".ts", ".tsx", ".json", ".yaml", ".yml", ".csv", ".html", ".css", ".log", ".toml", ".cpp", ".h"}:
        raw = path.read_bytes()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("gb18030")
        pages = [(None, text)]
    else:
        raise ValueError("目前支持 PDF、DOCX、XLSX、文本、代码和常见图片；旧版 DOC/XLS 请另存为新格式")
    if sum(len(text) for _, text in pages) > 1000000:
        raise ValueError("文本超过 100 万字符，请先拆分文件")
    chunks = []
    for page, text in pages:
        for start in range(0, len(text), 900):
            part = text[start:start + 1100].strip()
            if part:
                chunks.append({"source": str(path), "filename": path.name, "page": page,
                               "chunk": len(chunks) + 1, "text": part})
    if not chunks:
        raise ValueError("没有可读取的文字；扫描 PDF 需要先进行 OCR")
    return chunks


def _windows_ocr(path: Path) -> str:
    if sys.platform != "win32":
        raise ValueError("图片文字识别目前需要 Windows")
    quoted = str(path).replace("'", "''")
    script = """$ErrorActionPreference='Stop'
[Console]::OutputEncoding=[Text.Encoding]::UTF8
Add-Type -AssemblyName System.Runtime.WindowsRuntime
$null=[Windows.Storage.StorageFile,Windows.Storage,ContentType=WindowsRuntime]
$null=[Windows.Graphics.Imaging.BitmapDecoder,Windows.Foundation,ContentType=WindowsRuntime]
$null=[Windows.Media.Ocr.OcrEngine,Windows.Foundation,ContentType=WindowsRuntime]
function Await($operation,$type) {
 $method=[System.WindowsRuntimeSystemExtensions].GetMethods() | Where-Object {$_.Name -eq 'AsTask' -and $_.IsGenericMethod -and $_.GetParameters().Count -eq 1 -and $_.GetParameters()[0].ParameterType.Name -eq 'IAsyncOperation`1'} | Select-Object -First 1
 $task=$method.MakeGenericMethod($type).Invoke($null,@($operation));$task.Wait();$task.Result
}
""" + f"$file=Await ([Windows.Storage.StorageFile]::GetFileFromPathAsync('{quoted}')) ([Windows.Storage.StorageFile])\n" + """
$stream=Await ($file.OpenAsync([Windows.Storage.FileAccessMode]::Read)) ([Windows.Storage.Streams.IRandomAccessStream])
try {
 $decoder=Await ([Windows.Graphics.Imaging.BitmapDecoder]::CreateAsync($stream)) ([Windows.Graphics.Imaging.BitmapDecoder])
 $bitmap=Await ($decoder.GetSoftwareBitmapAsync()) ([Windows.Graphics.Imaging.SoftwareBitmap])
 $engine=[Windows.Media.Ocr.OcrEngine]::TryCreateFromUserProfileLanguages()
 if (!$engine) {throw 'Windows OCR language pack unavailable'}
 try { (Await ($engine.RecognizeAsync($bitmap)) ([Windows.Media.Ocr.OcrResult])).Text }
 finally {$bitmap.Dispose()}
} finally {$stream.Dispose()}
"""
    result = subprocess.run(["powershell.exe", "-NoProfile", "-NonInteractive", "-Command", script],
        capture_output=True, encoding="utf-8-sig", errors="replace", timeout=45,
        creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    if result.returncode:
        raise ValueError("Windows 图片识别失败，请检查系统 OCR 语言包或改用清晰 PNG 图片")
    return re.sub(r"(?<=[\u4e00-\u9fff])\s+(?=[\u4e00-\u9fff])", "", result.stdout.strip())


class PersonalContext:
    def __init__(self, path=None):
        self.path = Path(path or Path(config.CONFIG_DIR) / "context.sqlite3")
        self.path.parent.mkdir(parents=True, exist_ok=True)
        with self.connect() as db:
            db.executescript("""
                CREATE TABLE IF NOT EXISTS documents(id TEXT PRIMARY KEY, source TEXT, filename TEXT, digest TEXT UNIQUE, created_at TEXT);
                CREATE TABLE IF NOT EXISTS chunks(id INTEGER PRIMARY KEY, document_id TEXT REFERENCES documents(id) ON DELETE CASCADE, page TEXT, ordinal INTEGER, text TEXT);
                CREATE VIRTUAL TABLE IF NOT EXISTS knowledge_fts USING fts5(text, tokenize='trigram');
                CREATE TABLE IF NOT EXISTS memories(id TEXT PRIMARY KEY, topic TEXT UNIQUE, content TEXT, updated_at TEXT);
                PRAGMA user_version=1;
            """)

    @contextmanager
    def connect(self):
        db = sqlite3.connect(self.path, timeout=5)
        db.row_factory = sqlite3.Row
        db.execute("PRAGMA foreign_keys=ON")
        try:
            with db:
                yield db
        finally:
            db.close()

    def add_document(self, chunks):
        digest = hashlib.sha256(json.dumps(chunks, ensure_ascii=False).encode()).hexdigest()
        with self.connect() as db:
            existing = db.execute("SELECT id FROM documents WHERE digest=?", (digest,)).fetchone()
            if existing:
                return existing["id"]
            identity = uuid4().hex
            db.execute("INSERT INTO documents VALUES (?,?,?,?,?)", (identity, chunks[0]["source"], chunks[0]["filename"], digest, datetime.now(timezone.utc).isoformat()))
            for chunk in chunks:
                cursor = db.execute("INSERT INTO chunks(document_id,page,ordinal,text) VALUES (?,?,?,?)",
                    (identity, str(chunk["page"] or ""), chunk["chunk"], chunk["text"]))
                db.execute("INSERT INTO knowledge_fts(rowid,text) VALUES (?,?)", (cursor.lastrowid, chunk["text"]))
            return identity

    def search(self, query, limit=6):
        query = query.strip()[:300]
        if not query:
            return []
        with self.connect() as db:
            if len(query) >= 3:
                match = '"' + query.replace('"', '""') + '"'
                rows = db.execute("SELECT c.*,d.filename,d.source FROM knowledge_fts f JOIN chunks c ON c.id=f.rowid JOIN documents d ON d.id=c.document_id WHERE knowledge_fts MATCH ? ORDER BY rank LIMIT ?", (match, limit)).fetchall()
            else:
                rows = []
            if not rows:
                pattern = "%" + query.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_") + "%"
                rows = db.execute("SELECT c.*,d.filename,d.source FROM chunks c JOIN documents d ON d.id=c.document_id WHERE c.text LIKE ? ESCAPE '\\' LIMIT ?", (pattern, limit)).fetchall()
            return [dict(row) for row in rows]

    def memories(self):
        with self.connect() as db:
            return [dict(row) for row in db.execute("SELECT * FROM memories ORDER BY updated_at DESC LIMIT 100")]

    def remember(self, topic, content, replace=False):
        topic, content = topic.strip()[:80], content.strip()[:2000]
        if not topic or not content:
            raise ValueError("记忆主题和内容不能为空")
        with self.connect() as db:
            old = db.execute("SELECT * FROM memories WHERE topic=?", (topic,)).fetchone()
            if old and not replace and old["content"] != content:
                raise ValueError("同一主题已有不同记忆，请确认替换")
            identity = old["id"] if old else uuid4().hex
            db.execute("INSERT INTO memories VALUES (?,?,?,?) ON CONFLICT(topic) DO UPDATE SET content=excluded.content,updated_at=excluded.updated_at",
                       (identity, topic, content, datetime.now(timezone.utc).isoformat()))
            return identity

    def forget(self, identity):
        with self.connect() as db:
            db.execute("DELETE FROM memories WHERE id=?", (identity,))
